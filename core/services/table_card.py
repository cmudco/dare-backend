"""Search summaries for spreadsheet files.

Rows of a table carry no meaning a similarity search can match, and questions
about tables (averages, maxima, counts) need every row, so embedding the rows
gives confident answers built from a few random ones while flooding the index
(a 0.8 MB CSV became ~32k chunks). A spreadsheet is indexed as one card
instead: its shape, its columns and a few sample rows, enough for search to
find "the file with the PM1 data". Exact answers come from the Code sandbox,
which reads the whole file.
"""

import csv
import io
import math
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import PurePosixPath
from typing import Iterable, Iterator, List, Optional, Sequence, Tuple

from openpyxl import load_workbook

TABLE_EXTENSIONS = frozenset({"csv", "tsv", "xlsx", "xlsm", "xls"})

# Tables this small are indexed whole; their rows are few enough to matter.
FULL_ROWS_LIMIT = 30
SAMPLE_ROWS = 5
MAX_COLUMNS_LISTED = 150
MAX_SAMPLE_COLUMNS = 20
CELL_CHARS = 40
MISSING_MARKERS = frozenset(
    {"", "na", "n/a", "nan", "null", "none", "-", "#n/a", "inf", "-inf", "+inf"}
)


def is_table_file(filename: str) -> bool:
    return PurePosixPath(filename or "").suffix.lower().lstrip(".") in TABLE_EXTENSIONS


@dataclass
class _Column:
    name: str
    numbers: int = 0
    dates: int = 0
    texts: int = 0
    number_range: Optional[Tuple[float, float]] = None
    date_range: Optional[Tuple[date, date]] = None
    examples: List[str] = field(default_factory=list)

    def observe(self, value) -> None:
        if _is_missing(value):
            return
        number = _as_number(value)
        if number is not None:
            self.numbers += 1
            self.number_range = _widen(self.number_range, number)
        elif isinstance(value, (datetime, date)):
            self.dates += 1
            self.date_range = _widen(self.date_range, value)
        else:
            self.texts += 1
            text = _cell(value)
            if len(self.examples) < 3 and text not in self.examples:
                self.examples.append(text)

    def describe(self) -> str:
        seen = self.numbers + self.dates + self.texts
        if not seen:
            return f"- {self.name} (empty)"
        if self.texts == 0 and self.numbers >= self.dates:
            low, high = self.number_range
            return f"- {self.name} (number, {_cell(low)} to {_cell(high)})"
        if self.texts == 0:
            low, high = self.date_range
            return f"- {self.name} (date, {_cell(low)} to {_cell(high)})"
        examples = ", ".join(f'"{example}"' for example in self.examples)
        return f"- {self.name} (text, e.g. {examples})"


@dataclass
class _Sheet:
    title: Optional[str]
    header: Sequence[str]
    rows: int = 0
    samples: List[Sequence] = field(default_factory=list)
    columns: List[_Column] = field(default_factory=list)


def build_table_card(data: bytes, filename: str) -> str:
    """The text a spreadsheet is searched by: its shape, columns and samples."""
    sheets = [_summarize(title, rows) for title, rows in _read_sheets(data, filename)]
    sheets = [sheet for sheet in sheets if sheet.header]
    lines = [
        f"Table file: {PurePosixPath(filename).name}",
        "This is a summary of a spreadsheet for search. Its full rows are analyzed "
        "with the Code sandbox.",
    ]
    for sheet in sheets:
        lines.append("")
        prefix = f'Sheet "{sheet.title}": ' if sheet.title else ""
        lines.append(f"{prefix}{sheet.rows:,} rows × {len(sheet.header)} columns.")
        lines.append("Columns:")
        lines.extend(column.describe() for column in sheet.columns[:MAX_COLUMNS_LISTED])
        hidden = len(sheet.columns) - MAX_COLUMNS_LISTED
        if hidden > 0:
            lines.append(f"- … and {hidden} more columns")
        if sheet.samples:
            whole = sheet.rows <= FULL_ROWS_LIMIT
            lines.append("All rows:" if whole else f"First {len(sheet.samples)} rows:")
            lines.append(_row_line(sheet.header))
            lines.extend(_row_line(row) for row in sheet.samples)
    if not sheets:
        lines.append("")
        lines.append("The spreadsheet has no readable rows.")
    return "\n".join(lines)


def _summarize(title: Optional[str], rows: Iterator[Sequence]) -> _Sheet:
    header: Sequence = ()
    for row in rows:
        if any(_cell(value) for value in row):
            header = [_cell(value) or f"column {i + 1}" for i, value in enumerate(row)]
            break
    sheet = _Sheet(title=title, header=header)
    sheet.columns = [_Column(name) for name in header]
    sample_limit = FULL_ROWS_LIMIT + 1
    for row in rows:
        if not any(_cell(value) for value in row):
            continue
        sheet.rows += 1
        if len(sheet.samples) < sample_limit:
            sheet.samples.append(row)
        for column, value in zip(sheet.columns, row):
            column.observe(value)
    if sheet.rows > FULL_ROWS_LIMIT:
        sheet.samples = sheet.samples[:SAMPLE_ROWS]
    return sheet


def _read_sheets(
    data: bytes, filename: str
) -> Iterable[Tuple[Optional[str], Iterator[Sequence]]]:
    extension = PurePosixPath(filename).suffix.lower()
    if extension in (".xlsx", ".xlsm"):
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            for sheet in workbook.worksheets:
                yield sheet.title, sheet.iter_rows(values_only=True)
        finally:
            workbook.close()
    elif extension == ".xls":
        import xlrd

        book = xlrd.open_workbook(file_contents=data)
        for sheet in book.sheets():
            yield sheet.name, (sheet.row_values(i) for i in range(sheet.nrows))
    else:
        text = data.decode("utf-8-sig", errors="replace")
        delimiter = "\t" if extension == ".tsv" else _sniff_delimiter(text)
        yield None, csv.reader(io.StringIO(text), delimiter=delimiter)


def _sniff_delimiter(text: str) -> str:
    try:
        return csv.Sniffer().sniff(text[:20000], delimiters=",;\t|").delimiter
    except csv.Error:
        return ","


def _widen(bounds, value):
    if bounds is None:
        return value, value
    return min(bounds[0], value), max(bounds[1], value)


def _is_missing(value) -> bool:
    if value is None:
        return True
    if isinstance(value, str):
        return value.strip().lower() in MISSING_MARKERS
    return isinstance(value, float) and not math.isfinite(value)


def _as_number(value) -> Optional[float]:
    """A finite number, or None; Inf is a placeholder, not a value to range."""
    if isinstance(value, bool):
        return None
    if isinstance(value, str):
        try:
            value = float(value.replace(",", ""))
        except ValueError:
            return None
    if isinstance(value, (int, float)) and math.isfinite(value):
        return value
    return None


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and math.isfinite(value):
        value = int(value) if value.is_integer() else float(f"{value:.6g}")
    if isinstance(value, datetime) and not (value.hour or value.minute or value.second):
        value = value.date()
    text = " ".join(str(value).split())
    return text if len(text) <= CELL_CHARS else text[: CELL_CHARS - 1] + "…"


def _row_line(row: Sequence) -> str:
    cells = [_cell(value) for value in list(row)[:MAX_SAMPLE_COLUMNS]]
    extra = len(row) - MAX_SAMPLE_COLUMNS
    return " | ".join(cells) + (f" | … (+{extra} columns)" if extra > 0 else "")
